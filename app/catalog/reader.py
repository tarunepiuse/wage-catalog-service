"""Reads and structurally validates the payslip master-data input (XLSX workbook or CSV files)."""

import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

MASTER_SHEET = "master data"
LINES_SHEET = "pay line items"

MASTER_REQUIRED = frozenset({"payslip_id", "employee_id", "molga_country_grouping", "employment_group"})
# `wage_code` is expected in the current template but optional, so files in the original template still load
# (their codes then come from the mapping). Optional columns used when present: payroll_currency, applies_to.
LINES_REQUIRED = frozenset({"payslip_id", "employee_id", "section", "wage_type", "amount", "line_category"})
# A stand-alone line-items CSV must carry these master-data columns itself.
JOINED_EXTRA = frozenset({"molga_country_grouping", "employment_group"})

MAX_ROWS = 500_000
MAX_XLSX_UNCOMPRESSED = 200 * 1024 * 1024  # zip-bomb guard


class InputFormatError(ValueError):
    """The upload doesn't match the expected template. `errors` are user-facing, row-level messages."""

    def __init__(self, message: str, errors: list[str] | None = None):
        super().__init__(message)
        self.errors = errors or []


@dataclass(frozen=True)
class Row:
    number: int              # 1-based row number in the source sheet, for error messages
    values: dict[str, str]

    def __getitem__(self, column: str) -> str:
        return self.values.get(column, "")


@dataclass(frozen=True)
class Table:
    name: str
    columns: tuple[str, ...]
    rows: list[Row]

    def has(self, column: str) -> bool:
        return column in self.columns


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v).upper()
    if isinstance(v, float):
        # repr() is the shortest exact round-trip form, so 2707.66 stays "2707.66" (never 2707.6599999…);
        # formatting through Decimal avoids scientific notation such as 1e-05; 10.0 (a numeric molga) -> "10".
        text = format(Decimal(repr(v)), "f")
        return text.removesuffix(".0") if v.is_integer() else text
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _to_table(name: str, raw_rows, required: frozenset[str]) -> Table:
    it = iter(raw_rows)
    header = next(it, None)
    if header is None:
        raise InputFormatError(f"'{name}' is empty")
    columns = tuple(_cell_to_str(h).lower() for h in header)
    named = [c for c in columns if c]
    duplicates = sorted({c for c in named if named.count(c) > 1})
    if duplicates:
        raise InputFormatError(f"'{name}' has duplicate column(s): {', '.join(duplicates)}")
    missing = required - set(columns)
    if missing:
        raise InputFormatError(f"'{name}' is missing required column(s): {', '.join(sorted(missing))}")

    rows: list[Row] = []
    for number, raw in enumerate(it, start=2):
        values = [_cell_to_str(v) for v in raw]
        if not any(values):
            continue
        if len(rows) >= MAX_ROWS:
            raise InputFormatError(f"'{name}' exceeds the {MAX_ROWS:,} row limit")
        rows.append(Row(number, {c: v for c, v in zip(columns, values, strict=False) if c}))
    if not rows:
        raise InputFormatError(f"'{name}' has a header but no data rows")
    return Table(name, columns, rows)


def read_xlsx(path: Path) -> tuple[Table, Table]:
    if not zipfile.is_zipfile(path):
        raise InputFormatError("File has an .xlsx extension but is not a valid Excel workbook")
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist()) > MAX_XLSX_UNCOMPRESSED:
            raise InputFormatError("Workbook expands to an unreasonable size; refusing to process it")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        raise InputFormatError("File could not be opened as an .xlsx workbook") from None
    try:
        sheets = {ws.title.strip().lower(): ws for ws in wb.worksheets}
        if MASTER_SHEET not in sheets or LINES_SHEET not in sheets:
            raise InputFormatError("Workbook must contain sheets 'Master Data' and 'Pay Line Items' "
                                   f"(found: {', '.join(wb.sheetnames)})")
        tables = []
        for key, label, required in ((MASTER_SHEET, "Master Data", MASTER_REQUIRED),
                                     (LINES_SHEET, "Pay Line Items", LINES_REQUIRED)):
            ws = sheets[key]
            ws.reset_dimensions()  # some producers write a wrong <dimension>, which truncates read-only reads
            tables.append(_to_table(label, ws.iter_rows(values_only=True), required))
        return tables[0], tables[1]
    finally:
        wb.close()


def _csv_rows(path: Path, label: str):
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):  # cp1252 = Excel "Save as CSV" on Windows
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise InputFormatError(f"{label} is not UTF-8 or Windows-1252 encoded text")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return csv.reader(io.StringIO(text, newline=""), dialect)


def read_csv(lines_path: Path, master_path: Path | None) -> tuple[Table | None, Table]:
    if master_path is None:
        return None, _to_table("Pay Line Items CSV", _csv_rows(lines_path, "Line-items CSV"),
                               LINES_REQUIRED | JOINED_EXTRA)
    master = _to_table("Master Data CSV", _csv_rows(master_path, "Master-data CSV"), MASTER_REQUIRED)
    lines = _to_table("Pay Line Items CSV", _csv_rows(lines_path, "Line-items CSV"), LINES_REQUIRED)
    return master, lines
