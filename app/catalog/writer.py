"""Writes the catalog to XLSX in the Wage_Type_Catalog template layout."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.catalog.builder import CatalogResult

# (header / attribute, number_format, column width) — mirrors the reference Wage_Type_Catalog workbook.
COLUMNS = [
    ("wage_type_local_lang", "@", 33),
    ("wage_code", "@", 11),
    ("no_of_occurrence", "0", 18),
    ("total_pay", "#,##0.00", 11),
    ("average_pay", "#,##0.00", 13),
    ("category", "@", 23),
    ("sub_category", "@", 23),
    ("position", "0", 10),
    ("employment_groups", "@", 30),
    ("molga", "@", 7),
]
MONEY = {"total_pay", "average_pay"}

_EDGE = Side(style="thin", color="D7DAD6")
BORDER = Border(left=_EDGE, right=_EDGE, top=_EDGE, bottom=_EDGE)
HEADER_FONT = Font(name="Verdana", size=9, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1B2750")
HEADER_ALIGN = Alignment(vertical="center", wrap_text=True)
BODY_FONT = Font(name="Verdana", size=9, color="333333")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")  # highlights provisional wage codes


def _text(cell, value: str) -> None:
    cell.value = value
    cell.data_type = "s"  # force text: input like "=HYPERLINK(...)" must never become a live formula


def write_catalog(result: CatalogResult, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Wage Type Catalog"

    for col, (name, _, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=col)
        _text(c, name)
        c.font, c.fill, c.alignment, c.border = HEADER_FONT, HEADER_FILL, HEADER_ALIGN, BORDER
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 30

    for r, row in enumerate(result.rows, start=2):
        for col, (name, fmt, _) in enumerate(COLUMNS, start=1):
            value = getattr(row, name)
            c = ws.cell(row=r, column=col)
            if isinstance(value, str):
                _text(c, value)
            else:
                # Decimal -> float only at the very end: the value is already rounded to cents.
                c.value = float(value) if name in MONEY else value
            c.number_format, c.font, c.border = fmt, BODY_FONT, BORDER
        if row.provisional_code:
            ws.cell(row=r, column=2).fill = FLAG_FILL

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUMNS)).column_letter}{len(result.rows) + 1}"

    if result.warnings:
        notes = wb.create_sheet("Processing Notes")
        notes.column_dimensions["A"].width = 120
        h = notes.cell(row=1, column=1)
        _text(h, "Processing Notes")
        h.font, h.fill = HEADER_FONT, HEADER_FILL
        for i, warning in enumerate(result.warnings, start=2):
            c = notes.cell(row=i, column=1)
            _text(c, warning)
            c.font, c.alignment = BODY_FONT, Alignment(wrap_text=True, vertical="top")

    wb.save(path)
