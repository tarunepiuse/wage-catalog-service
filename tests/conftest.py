import csv
import io
import os
import shutil
import tempfile
from pathlib import Path

import openpyxl
import pytest

_tmp = Path(tempfile.mkdtemp(prefix="wtc-test-"))
os.environ.setdefault("WTC_JWT_SECRET", "test-secret-" + "x" * 40)
os.environ["WTC_USER_DB_PATH"] = str(_tmp / "users.db")
os.environ["WTC_JOB_DIR"] = str(_tmp / "jobs")
os.environ["WTC_LOG_LEVEL"] = "WARNING"

from fastapi.testclient import TestClient  # noqa: E402

from app.config import BASE_DIR, get_settings  # noqa: E402
from app.main import create_app  # noqa: E402

FIXTURES = Path(__file__).parent / "fixtures"
INPUT = FIXTURES / "Payslip_Demo_Master_Data_updated.xlsx"   # current template: Pay Line Items has wage_code
LEGACY_INPUT = FIXTURES / "Payslip_Demo_Master_Data.xlsx"    # original template: codes come from the mapping
EXPECTED = FIXTURES / "Wage_Type_Catalog_Demo.xlsx"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PASSWORD = "correct-horse-battery"


def make_client(tmp_path: Path, **overrides) -> TestClient:
    mapping = tmp_path / "wage_codes.csv"
    if not mapping.exists():
        shutil.copy(BASE_DIR / "data" / "wage_codes.csv", mapping)
    settings = get_settings().model_copy(update={
        "user_db_path": tmp_path / "users.db",
        "job_dir": tmp_path / "jobs",
        "wage_code_map_path": mapping,
        **overrides,
    })
    return TestClient(create_app(settings))


@pytest.fixture
def client(tmp_path):
    with make_client(tmp_path) as c:
        store = c.app.state.user_store
        store.create("alice", PASSWORD, "user")
        store.create("bob", PASSWORD, "user")
        store.create("root", PASSWORD, "admin")
        yield c


def login(client, username="alice", password=PASSWORD) -> dict:
    r = client.post("/v1/auth/token", data={"username": username, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def upload_xlsx(client, headers, path=INPUT):
    with open(path, "rb") as f:
        return client.post("/v1/catalog/jobs", headers=headers, files={"file": (path.name, f, XLSX_MIME)})


def upload_csv(client, headers, lines: bytes, master: bytes | None = None):
    files = {"file": ("lines.csv", lines, "text/csv")}
    if master is not None:
        files["master_data"] = ("master.csv", master, "text/csv")
    return client.post("/v1/catalog/jobs", headers=headers, files=files)


def sheet_csv(sheet: str, delimiter: str = ",", source: Path = INPUT) -> bytes:
    ws = openpyxl.load_workbook(source)[sheet]
    buf = io.StringIO()
    csv.writer(buf, delimiter=delimiter, lineterminator="\n").writerows(ws.iter_rows(values_only=True))
    return buf.getvalue().encode()


def with_extra_line(line: str) -> bytes:
    return sheet_csv("Pay Line Items") + line.encode() + b"\n"


def rows_of(data: bytes, sheet: str | None = None) -> list[tuple]:
    wb = openpyxl.load_workbook(io.BytesIO(data))
    return list((wb[sheet] if sheet else wb.active).iter_rows(values_only=True))


def expected_rows() -> list[tuple]:
    return list(openpyxl.load_workbook(EXPECTED).active.iter_rows(values_only=True))


def job_dirs(client) -> list[Path]:
    return list(client.app.state.settings.job_dir.iterdir())
