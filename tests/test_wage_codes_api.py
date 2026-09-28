"""Wage-code behaviour for the current template (Pay Line Items carries `wage_code`) and the legacy one."""

import io
from decimal import ROUND_HALF_UP, Decimal

import openpyxl

from tests.conftest import (
    INPUT,
    LEGACY_INPUT,
    expected_rows,
    login,
    make_client,
    rows_of,
    sheet_csv,
    upload_csv,
    upload_xlsx,
)

MASTER = sheet_csv("Master Data")
LINES = sheet_csv("Pay Line Items")
FIRST_PAYSLIP = "us-90088-2025-02-22-000000900004502,90088"


def _line(wage_type, code, amount="55.00", category="EARNING"):
    return (f"{FIRST_PAYSLIP},EARNING,CURRENT,{code},{wage_type},,,,,{amount},,{amount},"
            f"EMPLOYEE,TAXABLE,OTHER,{category}\n").encode()


def _run(client, lines: bytes):
    h = login(client)
    r = upload_csv(client, h, lines, MASTER)
    if r.status_code != 201:
        return r, None
    return r, rows_of(client.get(r.json()["download_url"], headers=h).content)


def test_output_matches_the_catalog_embedded_in_the_updated_workbook(client):
    # The updated input carries its own formula-driven 'Wage Type Catalog' sheet (Excel-cached values).
    embedded = list(openpyxl.load_workbook(INPUT, data_only=True)["Wage Type Catalog"].iter_rows(values_only=True))
    h = login(client)
    job = upload_xlsx(client, h).json()
    ours = rows_of(client.get(job["download_url"], headers=h).content)
    assert len(ours) == len(embedded)
    for got, want in zip(ours, embedded, strict=True):
        # Excel caches the unrounded average (e.g. 360.895); the catalog output is rounded half-up to cents.
        # (Float round() would give 360.89 here — exactly the error the Decimal arithmetic avoids.)
        assert got[:4] + got[5:] == want[:4] + want[5:]
        if isinstance(want[4], (int, float)):
            want_avg = Decimal(repr(want[4])).quantize(Decimal("0.01"), ROUND_HALF_UP)
            assert Decimal(repr(got[4])) == want_avg
        else:
            assert got[4] == want[4]


def test_input_codes_do_not_depend_on_the_mapping(tmp_path):
    (tmp_path / "wage_codes.csv").write_text("molga,wage_type,wage_code\n")  # empty mapping
    with make_client(tmp_path) as c:
        c.app.state.user_store.create("alice", "correct-horse-battery")
        h = login(c)
        job = upload_xlsx(c, h).json()
        assert job["summary"]["code_sources"] == {"input": 164, "mapping": 0, "provisional": 0}
        assert rows_of(c.get(job["download_url"], headers=h).content) == expected_rows()


def test_legacy_template_without_wage_code_still_works(client):
    h = login(client)
    r = upload_xlsx(client, h, LEGACY_INPUT)
    assert r.status_code == 201, r.text
    assert r.json()["summary"]["code_sources"] == {"input": 0, "mapping": 164, "provisional": 0}
    assert rows_of(client.get(r.json()["download_url"], headers=h).content) == expected_rows()


def test_input_code_wins_over_the_mapping(client):
    r, rows = _run(client, LINES.replace(b",CURRENT,1000,Regular Salary,", b",CURRENT,1001,Regular Salary,"))
    assert r.status_code == 201 and r.json()["warnings"] == []
    assert rows[1][:2] == ("Regular Salary", "1001") and all(row[1] != "1000" for row in rows[1:])


def test_one_code_for_two_wage_types_is_rejected(client):
    r, _ = _run(client, LINES + _line("Shift Premium", "1000"))
    assert r.status_code == 422
    expected = "wage_code '1000' (molga 10) is used for 'Shift Premium' but also for 'Regular Salary'"
    assert expected in r.json()["errors"][0]


def test_name_spelling_variants_share_a_code(client):
    r, rows = _run(client, LINES + _line("regular  SALARY", "1000", "10.00"))
    assert r.status_code == 201, r.text
    assert rows[1][:3] == ("Regular Salary", "1000", 11)


def test_one_wage_type_under_two_codes_is_two_rows_with_a_warning(client):
    r, rows = _run(client, LINES + _line("Regular Salary", "1005", "10.00"))
    assert r.status_code == 201
    assert any("'Regular Salary' (molga 10) appears under several wage codes (1000, 1005)" in w
               for w in r.json()["warnings"])
    assert [row[1] for row in rows[1:3]] == ["1000", "1005"]


def test_blank_code_falls_back_to_the_mapping_with_a_warning(client):
    r, rows = _run(client, LINES + _line("Regular Salary", "", "10.00"))
    assert r.status_code == 201
    assert r.json()["summary"]["code_sources"] == {"input": 164, "mapping": 1, "provisional": 0}
    assert any("1 line(s) had an empty wage_code" in w for w in r.json()["warnings"])
    assert rows[1][:3] == ("Regular Salary", "1000", 11)


def test_blank_unmapped_code_gets_provisional_code_clear_of_input_codes(client):
    r, rows = _run(client, LINES + _line("Night Premium", "1120") + _line("Shift Premium", ""))
    assert r.status_code == 201
    assert any("'Shift Premium' (molga 10) has no wage_code" in w and "1130" in w for w in r.json()["warnings"])
    pairs = [row[:2] for row in rows]
    assert ("Night Premium", "1120") in pairs and ("Shift Premium", "1130") in pairs


def test_invalid_wage_code_is_rejected(client):
    r, _ = _run(client, LINES + _line("Shift Premium", "10 00"))
    assert r.status_code == 422 and "wage_code '10 00' is not a valid code" in r.json()["errors"][0]


def test_numeric_wage_code_cells_are_read_as_codes(client, tmp_path):
    wb = openpyxl.load_workbook(INPUT)
    ws = wb["Pay Line Items"]
    col = [c.value for c in ws[1]].index("wage_code") + 1
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=col)
        cell.value, cell.number_format = int(cell.value), "0"  # as if typed as numbers in Excel
    path = tmp_path / "numeric_codes.xlsx"
    wb.save(path)
    h = login(client)
    job = upload_xlsx(client, h, path).json()
    assert rows_of(client.get(job["download_url"], headers=h).content) == expected_rows()


def test_other_sheets_in_the_workbook_are_ignored(client, tmp_path):
    wb = openpyxl.load_workbook(INPUT)
    wb["Wage Type Catalog"]["A2"] = "tampered"  # must not leak into the output
    path = tmp_path / "with_catalog.xlsx"
    wb.save(path)
    h = login(client)
    job = upload_xlsx(client, h, path).json()
    content = client.get(job["download_url"], headers=h).content
    assert rows_of(content) == expected_rows()
    assert openpyxl.load_workbook(io.BytesIO(content)).sheetnames == ["Wage Type Catalog"]
