import io
import time

import openpyxl
import pytest

from app.config import BASE_DIR
from tests.conftest import (
    INPUT,
    XLSX_MIME,
    expected_rows,
    job_dirs,
    login,
    make_client,
    rows_of,
    sheet_csv,
    upload_csv,
    upload_xlsx,
    with_extra_line,
)

MASTER = sheet_csv("Master Data")
LINES = sheet_csv("Pay Line Items")
FIRST_PAYSLIP = "us-90088-2025-02-22-000000900004502,90088"


def _line(wage_type="Shift Premium", amount="55.00", section="EARNING", category="EARNING", applies="CURRENT",
          code=""):
    """One Pay Line Items CSV row in the current template (wage_code is the 5th column)."""
    return (f"{FIRST_PAYSLIP},{section},{applies},{code},{wage_type},,,,,{amount},,{amount},"
            f"EMPLOYEE,TAXABLE,OTHER,{category}")


def _download(client, headers, job):
    return client.get(f"/v1/catalog/jobs/{job['job_id']}/download", headers=headers)


# ---------- happy paths ----------

def test_xlsx_round_trip_matches_reference_and_cleans_up(client):
    h = login(client)
    r = upload_xlsx(client, h)
    assert r.status_code == 201, r.text
    job = r.json()
    assert r.headers["location"].endswith(f"/v1/catalog/jobs/{job['job_id']}")
    assert job["summary"] == {"payslips": 10, "line_items": 164, "wage_types": 20, "provisional_codes": 0,
                              "molgas": ["10"], "currencies": {"10": "USD"},
                              "code_sources": {"input": 164, "mapping": 0, "provisional": 0}}
    assert job["warnings"] == []
    assert job["input_files"] == [{"name": INPUT.name, "bytes": INPUT.stat().st_size}]
    assert len(job_dirs(client)) == 1

    assert client.get(f"/v1/catalog/jobs/{job['job_id']}", headers=h).json()["job_id"] == job["job_id"]
    assert [j["job_id"] for j in client.get("/v1/catalog/jobs", headers=h).json()] == [job["job_id"]]

    d = _download(client, h, job)
    assert d.status_code == 200
    assert d.headers["content-type"] == XLSX_MIME
    assert d.headers["cache-control"] == "no-store"
    assert job["output_filename"] in d.headers["content-disposition"]
    assert rows_of(d.content) == expected_rows()

    assert job_dirs(client) == []
    again = _download(client, h, job)
    assert again.status_code == 404 and again.json()["code"] == "job_not_found"
    assert client.get("/v1/catalog/jobs", headers=h).json() == []


def test_csv_pair_matches_reference(client):
    h = login(client)
    r = upload_csv(client, h, LINES, MASTER)
    assert r.status_code == 201, r.text
    assert rows_of(_download(client, h, r.json()).content) == expected_rows()


def test_prejoined_single_csv_matches_reference(client):
    wb = openpyxl.load_workbook(INPUT)
    header = [c.value for c in wb["Master Data"][1]]
    master = {row[1]: row for row in wb["Master Data"].iter_rows(min_row=2, values_only=True)}
    mi, gi = header.index("molga_country_grouping"), header.index("employment_group")
    out = []
    for i, row in enumerate(wb["Pay Line Items"].iter_rows(values_only=True)):
        extra = ["molga_country_grouping", "employment_group"] if i == 0 else [master[row[0]][mi], master[row[0]][gi]]
        out.append(",".join("" if v is None else str(v) for v in list(row) + extra))
    h = login(client)
    r = upload_csv(client, h, ("\n".join(out) + "\n").encode())
    assert r.status_code == 201, r.text
    assert rows_of(_download(client, h, r.json()).content) == expected_rows()


def test_semicolon_csv_is_detected(client):
    h = login(client)
    r = upload_csv(client, h, sheet_csv("Pay Line Items", ";"), sheet_csv("Master Data", ";"))
    assert r.status_code == 201, r.text
    assert rows_of(_download(client, h, r.json()).content) == expected_rows()


# ---------- wage codes ----------

def test_unmapped_wage_type_gets_flagged_provisional_code(client):
    h = login(client)
    r = upload_csv(client, h, with_extra_line(_line()), MASTER)
    assert r.status_code == 201, r.text
    job = r.json()
    assert job["summary"]["provisional_codes"] == 1
    assert any("Shift Premium" in w and "1120" in w for w in job["warnings"])
    content = _download(client, h, job).content
    assert ("Shift Premium", "1120", 1, 55.0, 55.0, "Earning", "EARNING", 4,
            "BW1-Biweekly Salaried Exempt", "10") in rows_of(content)
    assert "Processing Notes" in openpyxl.load_workbook(io.BytesIO(content)).sheetnames


def test_provisional_code_never_reuses_a_reserved_mapping_code(tmp_path):
    mapping = (BASE_DIR / "data" / "wage_codes.csv").read_text()
    (tmp_path / "wage_codes.csv").write_text(mapping + "*,Night Premium,1120\n")
    with make_client(tmp_path) as c:
        c.app.state.user_store.create("alice", "correct-horse-battery")
        h = login(c)
        r = upload_csv(c, h, with_extra_line(_line("Shift Premium")), MASTER)
        assert any("provisional code 1130" in w for w in r.json()["warnings"]), r.json()["warnings"]


def test_broken_mapping_returns_503_configuration_error(client):
    client.app.state.settings.wage_code_map_path.write_text("name,code\nx,1\n")
    r = upload_xlsx(client, login(client))
    assert r.status_code == 503 and r.json()["code"] == "configuration_error"
    assert job_dirs(client) == []
    assert client.get("/health/ready").status_code == 503


# ---------- payroll data validation ----------

def test_invalid_rows_are_reported_with_row_numbers(client):
    bad = with_extra_line(_line("Regular Salary", "abc").replace(FIRST_PAYSLIP, "unknown-payslip,00000"))
    r = upload_csv(client, login(client), bad, MASTER)
    assert r.status_code == 422
    body = r.json()
    assert r.headers["content-type"] == "application/problem+json"
    assert body["code"] == "invalid_input" and body["request_id"]
    assert body["errors"] == ["Pay Line Items CSV row 166: payslip_id 'unknown-payslip' not found in Master Data CSV"]
    assert job_dirs(client) == []


@pytest.mark.parametrize("amount", ["1234,56", "1.234,56", "1,234", "12a"])
def test_ambiguous_amounts_are_rejected_not_guessed(client, amount):
    r = upload_csv(client, login(client), with_extra_line(_line(amount=f'"{amount}"')), MASTER)
    assert r.status_code == 422, r.text
    assert "not an unambiguous number" in r.json()["errors"][0]


def test_mixed_currency_in_one_molga_is_rejected(client):
    master = MASTER.replace(b"USD,USD,BIWEEKLY,2025-02-22", b"EUR,EUR,BIWEEKLY,2025-02-22", 1)
    r = upload_csv(client, login(client), LINES, master)
    assert r.status_code == 422
    assert "mixes payroll currencies (EUR, USD)" in r.json()["errors"][0]


def test_non_current_lines_are_included_but_warned(client):
    r = upload_csv(client, login(client), with_extra_line(_line("Regular Salary", "10.00", applies="RETRO")), MASTER)
    assert r.status_code == 201
    assert any("applies_to other than CURRENT (RETRO: 1)" in w for w in r.json()["warnings"])


def test_formula_like_text_is_written_as_text(client):
    h = login(client)
    r = upload_csv(client, h, with_extra_line(_line('"=HYPERLINK(""http://x"")"', "1")), MASTER)
    ws = openpyxl.load_workbook(io.BytesIO(_download(client, h, r.json()).content)).active
    cell = next(c for c in ws["A"] if str(c.value).startswith("="))
    assert cell.data_type == "s"


# ---------- request validation ----------

def test_rejects_unsupported_extension(client):
    r = client.post("/v1/catalog/jobs", headers=login(client), files={"file": ("x.pdf", b"%PDF", "application/pdf")})
    assert r.status_code == 415 and r.json()["code"] == "unsupported_file_type"


def test_rejects_fake_xlsx(client):
    r = client.post("/v1/catalog/jobs", headers=login(client), files={"file": ("x.xlsx", b"not a zip", XLSX_MIME)})
    assert r.status_code == 422
    assert job_dirs(client) == []


def test_rejects_workbook_without_expected_sheets(client, tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Other"
    wb.save(tmp_path / "x.xlsx")
    r = upload_xlsx(client, login(client), tmp_path / "x.xlsx")
    assert r.status_code == 422 and "Master Data" in r.json()["detail"]


def test_rejects_single_csv_without_master_columns(client):
    r = upload_csv(client, login(client), LINES)
    assert r.status_code == 422 and "employment_group" in r.json()["detail"]


def test_rejects_missing_file_field(client):
    r = client.post("/v1/catalog/jobs", headers=login(client), data={"x": "1"})
    assert r.status_code == 422 and r.json()["code"] == "missing_file"


def test_rejects_master_data_with_xlsx(client):
    with open(INPUT, "rb") as f:
        r = client.post("/v1/catalog/jobs", headers=login(client),
                        files={"file": (INPUT.name, f, XLSX_MIME), "master_data": ("m.csv", MASTER, "text/csv")})
    assert r.status_code == 422


# ---------- limits ----------

def test_per_file_size_limit(client):
    client.app.state.settings = client.app.state.settings.model_copy(update={"max_upload_mb": 1})
    r = upload_csv(client, login(client), LINES + b"x" * (1024 * 1024), MASTER)
    assert r.status_code == 413 and r.json()["code"] == "file_too_large"
    assert job_dirs(client) == []


def test_chunked_body_over_limit_is_cut_off(tmp_path):
    with make_client(tmp_path, max_upload_mb=1) as c:
        c.app.state.user_store.create("alice", "correct-horse-battery")
        h = login(c)

        def body():  # generator => chunked transfer, no Content-Length to check up front
            yield b"--b\r\nContent-Disposition: form-data; name=\"file\"; filename=\"a.csv\"\r\n\r\n"
            for _ in range(4):
                yield b"x" * (1024 * 1024)

        r = c.post("/v1/catalog/jobs", headers={**h, "content-type": "multipart/form-data; boundary=b"},
                   content=body())
        assert r.status_code == 413, r.text
        assert job_dirs(c) == []


def test_unauthenticated_upload_is_refused_before_the_body_is_read(client):
    # Larger than the request limit, yet the answer is 401, not 413: auth runs before the body is touched.
    client.app.state.settings = client.app.state.settings.model_copy(update={"max_upload_mb": 1})

    def body():
        yield b"--b\r\n"
        for _ in range(30):
            yield b"x" * (1024 * 1024)

    r = client.post("/v1/catalog/jobs", headers={"content-type": "multipart/form-data; boundary=b"}, content=body())
    assert r.status_code == 401


def test_pending_job_quota(tmp_path):
    with make_client(tmp_path, max_pending_jobs_per_user=2) as c:
        c.app.state.user_store.create("alice", "correct-horse-battery")
        h = login(c)
        assert upload_xlsx(c, h).status_code == 201
        assert upload_xlsx(c, h).status_code == 201
        r = upload_xlsx(c, h)
        assert r.status_code == 429 and r.json()["code"] == "too_many_pending_jobs"


# ---------- lifecycle ----------

def test_range_request_gets_the_whole_file_and_is_one_time(client):
    h = login(client)
    job = upload_xlsx(client, h).json()
    d = client.get(f"/v1/catalog/jobs/{job['job_id']}/download", headers={**h, "Range": "bytes=0-0"})
    assert d.status_code == 200 and d.headers["accept-ranges"] == "none"
    assert rows_of(d.content) == expected_rows()


def test_other_user_cannot_see_download_or_discard(client):
    job = upload_xlsx(client, login(client, "alice")).json()
    bob = login(client, "bob")
    for call in (client.get(f"/v1/catalog/jobs/{job['job_id']}", headers=bob),
                 client.get(f"/v1/catalog/jobs/{job['job_id']}/download", headers=bob),
                 client.delete(f"/v1/catalog/jobs/{job['job_id']}", headers=bob)):
        assert call.status_code == 404
    assert client.get("/v1/catalog/jobs", headers=bob).json() == []
    assert _download(client, login(client, "alice"), job).status_code == 200


def test_discard_deletes_files(client):
    h = login(client)
    job = upload_xlsx(client, h).json()
    assert client.delete(f"/v1/catalog/jobs/{job['job_id']}", headers=h).status_code == 204
    assert job_dirs(client) == []


def test_expired_job_is_unreachable_and_swept(client):
    h = login(client)
    job = upload_xlsx(client, h).json()
    store = client.app.state.job_store
    meta = store.get(job["job_id"], "alice")
    meta.expires_at = time.time() - 1
    store.commit(meta)
    assert _download(client, h, job).status_code == 404
    assert store.sweep() == [job["job_id"]]
    assert job_dirs(client) == []


def test_released_claim_can_be_downloaded_again(client):
    h = login(client)
    job = upload_xlsx(client, h).json()
    store = client.app.state.job_store
    meta, claimed = store.claim(job["job_id"], "alice")
    assert _download(client, h, job).status_code == 404  # claimed: a concurrent second download loses
    assert store.release(meta, claimed)
    assert _download(client, h, job).status_code == 200
